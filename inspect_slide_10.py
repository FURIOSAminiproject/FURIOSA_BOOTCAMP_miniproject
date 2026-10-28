import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pptx import Presentation

ppt_path = r"c:\furiosa-miniproject\FURIOSA_BOOTCAMP_miniproject\퓨리오사AIagent스쿨_미니프로젝트_IamyourJ_final_이지석_장동율.pptx"
prs = Presentation(ppt_path)
s = prs.slides[10]

for idx, shp in enumerate(s.shapes):
    txt = shp.text_frame.text if shp.has_text_frame else ""
    print(f"Shape {idx}: name={shp.name}, type={shp.shape_type}, pos=({shp.left}, {shp.top}, {shp.width}, {shp.height})")
    if txt:
        print(f"   text={repr(txt[:100])}")
    if shp.shape_type == 6: # group
        for cidx, c in enumerate(shp.shapes):
            ctxt = c.text_frame.text if c.has_text_frame else ""
            print(f"   child {cidx}: name={c.name}, type={c.shape_type}, pos=({c.left}, {c.top}, {c.width}, {c.height}), text={repr(ctxt[:60])}")
