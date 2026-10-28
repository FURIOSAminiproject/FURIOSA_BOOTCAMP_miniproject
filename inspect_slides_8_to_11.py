import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pptx import Presentation

ppt_path = r"c:\furiosa-miniproject\FURIOSA_BOOTCAMP_miniproject\퓨리오사AIagent스쿨_미니프로젝트_IamyourJ_final_이지석_장동율.pptx"
prs = Presentation(ppt_path)

for slide_idx in [8, 9, 10, 11]:
    slide = prs.slides[slide_idx]
    print(f"\n================ SLIDE {slide_idx} ================")
    for s in slide.shapes:
        if s.has_text_frame:
            text = "\n".join(p.text for p in s.text_frame.paragraphs if p.text.strip())
            if text:
                print(f"[{s.name}] ({s.left}, {s.top}, {s.width}, {s.height}):\n{text}\n")
