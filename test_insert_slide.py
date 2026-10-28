import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pptx import Presentation

ppt_path = r"c:\furiosa-miniproject\FURIOSA_BOOTCAMP_miniproject\퓨리오사AIagent스쿨_미니프로젝트_IamyourJ_final_이지석_장동율.pptx"
prs = Presentation(ppt_path)
init_len = len(prs.slides)
print(f"Initial slides count: {init_len}")

blank_layout = prs.slide_layouts[6]
new_slide = prs.slides.add_slide(blank_layout)
sldId = prs.slides._sldIdLst[-1]
prs.slides._sldIdLst.remove(sldId)
prs.slides._sldIdLst.insert(8, sldId)

print(f"New slides count: {len(prs.slides)}")
print(f"Slide at 7: {[p.text for s in prs.slides[7].shapes if s.has_text_frame for p in s.text_frame.paragraphs if p.text.strip()][:2]}")
print(f"Slide at 8: {[p.text for s in prs.slides[8].shapes if s.has_text_frame for p in s.text_frame.paragraphs if p.text.strip()][:2]}")
print(f"Slide at 9: {[p.text for s in prs.slides[9].shapes if s.has_text_frame for p in s.text_frame.paragraphs if p.text.strip()][:2]}")

# Clean up test by removing the inserted slide
prs.slides._sldIdLst.remove(sldId)
print(f"Cleaned up slides count: {len(prs.slides)}")
