# -*- coding: utf-8 -*-
"""
build_clean_architecture_slides.py
Slide 12(조장 날씨 예측 아키텍처)와 100% 동일한 네이티브 도형/배치/폰트/컬러 적용
모든 줄바꿈은 별도 paragraph로 생성하여 네모(□) 박스 문자 완전 제거
"""

import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

PPT_PATH = r"c:\furiosa-miniproject\FURIOSA_BOOTCAMP_miniproject\퓨리오사AIagent스쿨_미니프로젝트_IamyourJ_final_이지석_장동율.pptx"

prs = Presentation(PPT_PATH)

# 컬러 팔레트 (Slide 12와 100% 일치)
COLOR_MINT_LIGHT  = RGBColor(221, 248, 242) # #DDF8F2
COLOR_MINT_DARK   = RGBColor(162, 236, 220) # #A2ECDC
COLOR_WHITE       = RGBColor(255, 255, 255) # #FFFFFF
COLOR_BORDER_GRAY = RGBColor(148, 163, 184) # #94A3B8
COLOR_BORDER_MINT = RGBColor(94, 234, 212)  # #5EEAD4
COLOR_BORDER_TEAL = RGBColor(45, 212, 191)  # #2DD4BF
COLOR_TEXT_MAIN   = RGBColor(15, 23, 42)    # #0F172A
COLOR_TEXT_TITLE  = RGBColor(34, 59, 107)   # #223B6B
COLOR_TEXT_MUTED  = RGBColor(71, 85, 105)   # #475569
COLOR_ARROW       = RGBColor(13, 148, 136)  # #0D9488
COLOR_LINE_GRAY   = RGBColor(203, 213, 225) # #CBD5E1

def clean_str(s):
    if not s:
        return ""
    return str(s).replace('\r', '').replace('\n', '').replace('\x0b', '').strip()

def setup_para(p, text, font_size=15, bold=True, color=COLOR_TEXT_MAIN, align=PP_ALIGN.CENTER):
    p.text = clean_str(text)
    p.font.name = "Malgun Gothic"
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align

def add_clean_box(slide, left, top, width, height, text_items, fill_color, border_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
        
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.08)
    tf.margin_bottom = Inches(0.08)
    
    is_first = True
    for item in text_items:
        line, is_bold, sz, col = item
        clean_text = clean_str(line)
        if not clean_text:
            p = tf.paragraphs[0] if is_first else tf.add_paragraph()
            is_first = False
            setup_para(p, "", font_size=4, bold=False, color=col)
            continue
        p = tf.paragraphs[0] if is_first else tf.add_paragraph()
        is_first = False
        setup_para(p, clean_text, font_size=sz, bold=is_bold, color=col, align=PP_ALIGN.CENTER)
    return shape

def add_arrow(slide, left, top, width=300000, height=440000, color=COLOR_ARROW):
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = color
    arrow.line.fill.background()
    return arrow

def setup_stage_label(slide, left, top, width, height, stage_title, stage_desc):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    setup_para(p0, clean_str(stage_title), font_size=17, bold=True, color=COLOR_TEXT_MAIN, align=PP_ALIGN.CENTER)
    p1 = tf.add_paragraph()
    setup_para(p1, clean_str(stage_desc), font_size=14, bold=False, color=COLOR_TEXT_MUTED, align=PP_ALIGN.CENTER)
    return tb

# =============================================================================
# 1. SLIDE 7: BackEnd – 여행 경로 생성
# =============================================================================
slide7 = prs.slides[7]
for s in list(slide7.shapes):
    if s.name not in ["슬라이드 번호 개체 틀 1", "TextBox 2", "직선 연결선 3"]:
        sp = s._element
        sp.getparent().remove(sp)

for s in slide7.shapes:
    if s.name == "TextBox 2":
        s.width = 6500000
        setup_para(s.text_frame.paragraphs[0], "BackEnd – 여행 경로 생성", font_size=28, bold=True, color=COLOR_TEXT_TITLE, align=PP_ALIGN.LEFT)
    elif s.name == "슬라이드 번호 개체 틀 1":
        s.text_frame.text = "7"

sub7 = slide7.shapes.add_textbox(664358, 1197015, 6000000, 494494)
setup_para(sub7.text_frame.paragraphs[0], "Multi-stage Routing & Spot Curation Architecture", font_size=18, bold=True, color=COLOR_TEXT_MUTED, align=PP_ALIGN.LEFT)

line7 = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, 5050000, 1950000, Pt(1.5), 3450000)
line7.fill.solid()
line7.fill.fore_color.rgb = COLOR_LINE_GRAY
line7.line.fill.background()

# Stage 1 (좌측)
add_clean_box(
    slide7, 500000, 2180000, 580000, 3150000,
    [
        ("출", True, 16, COLOR_TEXT_MAIN),
        ("발", True, 16, COLOR_TEXT_MAIN),
        ("지", True, 16, COLOR_TEXT_MAIN),
        ("/", True, 14, COLOR_TEXT_MUTED),
        ("도", True, 16, COLOR_TEXT_MAIN),
        ("착", True, 16, COLOR_TEXT_MAIN),
        ("지", True, 16, COLOR_TEXT_MAIN),
    ],
    COLOR_WHITE, COLOR_BORDER_GRAY
)
add_arrow(slide7, 1140000, 3530000)

add_clean_box(
    slide7, 1500000, 2180000, 2050000, 1460000,
    [
        ("PostgreSQL 24h 캐시", True, 16, RGBColor(6, 78, 59)),
        ("", False, 4, COLOR_TEXT_MAIN),
        ("(routes_cache / SHA-256)", False, 13, COLOR_TEXT_MUTED),
        ("10ms 즉시 응답 (호출 85% 절감)", False, 12, RGBColor(4, 120, 87)),
    ],
    COLOR_MINT_LIGHT, COLOR_BORDER_MINT
)

add_clean_box(
    slide7, 1500000, 3870000, 2050000, 1460000,
    [
        ("OpenRouteService", True, 16, RGBColor(17, 94, 89)),
        ("", False, 4, COLOR_TEXT_MAIN),
        ("(자차 / 도보 경로 엔진)", False, 13, COLOR_TEXT_MUTED),
        ("권장 및 최단 대안 경로 산출", False, 12, RGBColor(13, 148, 136)),
    ],
    COLOR_MINT_DARK, COLOR_BORDER_TEAL
)
add_arrow(slide7, 3610000, 3530000)

add_clean_box(
    slide7, 3970000, 2180000, 950000, 3150000,
    [
        ("GeoJSON", True, 15.5, COLOR_TEXT_MAIN),
        ("LineString", True, 15.5, COLOR_TEXT_MAIN),
        ("", False, 8, COLOR_TEXT_MAIN),
        ("(도로 궤적)", False, 13, COLOR_TEXT_MUTED),
        ("좌표열", False, 13, COLOR_TEXT_MUTED),
    ],
    COLOR_WHITE, COLOR_BORDER_GRAY
)

setup_stage_label(slide7, 1100000, 5450000, 3200000, 882293, "Stage 1", "이동 경로 생성 및 24h 캐시 탐색")

# Stage 2 (우측)
add_arrow(slide7, 4980000, 3530000)

add_clean_box(
    slide7, 5360000, 2180000, 2200000, 1460000,
    [
        ("타임라인 선형 보간", True, 16, RGBColor(6, 78, 59)),
        ("", False, 4, COLOR_TEXT_MAIN),
        ("(Linear Interpolation)", False, 13, COLOR_TEXT_MUTED),
        ("주행 시간대별 통과 지점 역산", False, 12, RGBColor(4, 120, 87)),
    ],
    COLOR_MINT_LIGHT, COLOR_BORDER_MINT
)

add_clean_box(
    slide7, 5360000, 3870000, 2200000, 1460000,
    [
        ("PostGIS 5km 공간 질의", True, 16, RGBColor(17, 94, 89)),
        ("", False, 4, COLOR_TEXT_MAIN),
        ("(ST_DWithin 수직거리)", False, 13, COLOR_TEXT_MUTED),
        ("도로변 5km 이내 POI 30ms 추출", False, 12, RGBColor(13, 148, 136)),
    ],
    COLOR_MINT_DARK, COLOR_BORDER_TEAL
)
add_arrow(slide7, 7620000, 3530000)

add_clean_box(
    slide7, 7980000, 2180000, 2250000, 3150000,
    [
        ("TourAPI 4.0", True, 17, RGBColor(6, 78, 59)),
        ("계층 큐레이션", True, 17, RGBColor(6, 78, 59)),
        ("", False, 10, COLOR_TEXT_MAIN),
        ("관광·숙박·맛집 3단계 필터링", False, 13, COLOR_TEXT_MAIN),
        ("", False, 4, COLOR_TEXT_MAIN),
        ("여행자 취향 키워드 매핑", False, 13, COLOR_TEXT_MAIN),
        ("", False, 4, COLOR_TEXT_MAIN),
        ("고화질 이미지 & 주소 결합", False, 13, COLOR_TEXT_MAIN),
    ],
    COLOR_MINT_LIGHT, COLOR_BORDER_MINT
)
add_arrow(slide7, 10290000, 3530000)

add_clean_box(
    slide7, 10650000, 2180000, 950000, 3150000,
    [
        ("Leaflet", True, 16, COLOR_TEXT_MAIN),
        ("여행", True, 16, COLOR_TEXT_MAIN),
        ("일정", True, 16, COLOR_TEXT_MAIN),
        ("표출", True, 16, COLOR_TEXT_MAIN),
        ("", False, 8, COLOR_TEXT_MAIN),
        ("(지도 시각화)", False, 12.5, COLOR_TEXT_MUTED),
    ],
    COLOR_WHITE, COLOR_BORDER_GRAY
)

setup_stage_label(slide7, 7000000, 5450000, 4000000, 882293, "Stage 2", "선형 보간 타임라인 및 길목 명소 큐레이션")


# =============================================================================
# 2. SLIDE 15: CI/CD
# =============================================================================
slide15 = prs.slides[15]
for s in list(slide15.shapes):
    if s.name not in ["슬라이드 번호 개체 틀 1", "TextBox 3", "직선 연결선 2"]:
        sp = s._element
        sp.getparent().remove(sp)

for s in slide15.shapes:
    if s.name == "TextBox 3":
        s.width = 6500000
        setup_para(s.text_frame.paragraphs[0], "CI/CD", font_size=28, bold=True, color=COLOR_TEXT_TITLE, align=PP_ALIGN.LEFT)
    elif s.name == "슬라이드 번호 개체 틀 1":
        s.text_frame.text = "15"

sub15 = slide15.shapes.add_textbox(664358, 1197015, 6500000, 494494)
setup_para(sub15.text_frame.paragraphs[0], "Two-stage Continuous Integration & Deployment Architecture", font_size=18, bold=True, color=COLOR_TEXT_MUTED, align=PP_ALIGN.LEFT)

line15 = slide15.shapes.add_shape(MSO_SHAPE.RECTANGLE, 5050000, 1950000, Pt(1.5), 3450000)
line15.fill.solid()
line15.fill.fore_color.rgb = COLOR_LINE_GRAY
line15.line.fill.background()

# Stage 1 (좌측)
add_clean_box(
    slide15, 500000, 2180000, 580000, 3150000,
    [
        ("코", True, 16, COLOR_TEXT_MAIN),
        ("드", True, 16, COLOR_TEXT_MAIN),
        ("", False, 6, COLOR_TEXT_MAIN),
        ("푸", True, 16, COLOR_TEXT_MAIN),
        ("시", True, 16, COLOR_TEXT_MAIN),
        ("", False, 6, COLOR_TEXT_MAIN),
        ("(main)", False, 13, COLOR_TEXT_MUTED),
    ],
    COLOR_WHITE, COLOR_BORDER_GRAY
)
add_arrow(slide15, 1140000, 3530000)

add_clean_box(
    slide15, 1500000, 2180000, 2050000, 1460000,
    [
        ("GitHub Cloud Runner", True, 16, RGBColor(6, 78, 59)),
        ("", False, 4, COLOR_TEXT_MAIN),
        ("(ubuntu-latest)", False, 13, COLOR_TEXT_MUTED),
        ("EC2 서버 부하 0% 사전 격리", False, 12, RGBColor(4, 120, 87)),
    ],
    COLOR_MINT_LIGHT, COLOR_BORDER_MINT
)

add_clean_box(
    slide15, 1500000, 3870000, 2050000, 1460000,
    [
        ("Vite 빌드 & tsc 검증", True, 16, RGBColor(17, 94, 89)),
        ("", False, 4, COLOR_TEXT_MAIN),
        ("(Fail-Fast 조기 방어)", False, 13, COLOR_TEXT_MUTED),
        ("에러 시 즉시 중단 & PR 알림", False, 12, RGBColor(13, 148, 136)),
    ],
    COLOR_MINT_DARK, COLOR_BORDER_TEAL
)
add_arrow(slide15, 3610000, 3530000)

add_clean_box(
    slide15, 3970000, 2180000, 950000, 3150000,
    [
        ("dist/", True, 16, COLOR_TEXT_MAIN),
        ("빌드", True, 16, COLOR_TEXT_MAIN),
        ("아티팩트", True, 16, COLOR_TEXT_MAIN),
        ("", False, 8, COLOR_TEXT_MAIN),
        ("(upload-artifact)", False, 12, COLOR_TEXT_MUTED),
    ],
    COLOR_WHITE, COLOR_BORDER_GRAY
)

setup_stage_label(slide15, 1100000, 5450000, 3200000, 882293, "Stage 1", "GitHub Cloud 사전 격리 빌드 (CI)")

# Stage 2 (우측)
add_arrow(slide15, 4980000, 3530000)

add_clean_box(
    slide15, 5360000, 2180000, 2200000, 1460000,
    [
        ("EC2 Self-Hosted Runner", True, 16, RGBColor(6, 78, 59)),
        ("", False, 4, COLOR_TEXT_MAIN),
        ("(SSH 미개방 보안 수신)", False, 13, COLOR_TEXT_MUTED),
        ("인바운드 22번 포트 차단 통신", False, 12, RGBColor(4, 120, 87)),
    ],
    COLOR_MINT_LIGHT, COLOR_BORDER_MINT
)

add_clean_box(
    slide15, 5360000, 3870000, 2200000, 1460000,
    [
        ("사전 백업 & 핫스왑", True, 16, RGBColor(17, 94, 89)),
        ("", False, 4, COLOR_TEXT_MAIN),
        ("(backups/ & --no-deps)", False, 13, COLOR_TEXT_MUTED),
        ("압축 백업 및 1초 롤백 보장", False, 12, RGBColor(13, 148, 136)),
    ],
    COLOR_MINT_DARK, COLOR_BORDER_TEAL
)
add_arrow(slide15, 7620000, 3530000)

add_clean_box(
    slide15, 7980000, 2180000, 2250000, 3150000,
    [
        ("Docker Compose", True, 17, RGBColor(6, 78, 59)),
        ("& 자가 치유", True, 17, RGBColor(6, 78, 59)),
        ("", False, 10, COLOR_TEXT_MAIN),
        ("/api/health 5회 헬스체크", False, 13, COLOR_TEXT_MAIN),
        ("", False, 4, COLOR_TEXT_MAIN),
        ("restart: always 3초 부활", False, 13, COLOR_TEXT_MAIN),
        ("", False, 4, COLOR_TEXT_MAIN),
        ("PostgreSQL 호스트 DB 격리", False, 13, COLOR_TEXT_MAIN),
    ],
    COLOR_MINT_LIGHT, COLOR_BORDER_MINT
)
add_arrow(slide15, 10290000, 3530000)

add_clean_box(
    slide15, 10650000, 2180000, 950000, 3150000,
    [
        ("무중단", True, 16, COLOR_TEXT_MAIN),
        ("서비스", True, 16, COLOR_TEXT_MAIN),
        ("배포", True, 16, COLOR_TEXT_MAIN),
        ("완료", True, 16, COLOR_TEXT_MAIN),
        ("", False, 8, COLOR_TEXT_MAIN),
        ("(Zero-Downtime)", False, 12, COLOR_TEXT_MUTED),
    ],
    COLOR_WHITE, COLOR_BORDER_GRAY
)

setup_stage_label(slide15, 7000000, 5450000, 4000000, 882293, "Stage 2", "EC2 인하우스 러너 무중단 롤아웃 (CD)")

prs.save(PPT_PATH)
print("Saved cleanly with zero newline characters!")
