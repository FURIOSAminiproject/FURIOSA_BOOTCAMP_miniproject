# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pptx import Presentation

prs = Presentation(r'c:\furiosa-miniproject\FURIOSA_BOOTCAMP_miniproject\퓨리오사AIagent스쿨_미니프로젝트_IamyourJ_final_이지석_장동율.pptx')

for idx in [9, 12]:
    slide = prs.slides[idx]
    print(f"\n==================== DETAIL SLIDE {idx} ====================")
    for s_idx, s in enumerate(slide.shapes):
        t = s.text.replace('\n', ' // ') if s.has_text_frame else ''
        fill = "None"
        line = "None"
        try:
            if s.fill.type == 1:
                fill = f"RGB({s.fill.fore_color.rgb})"
            elif s.fill.type == 0:
                fill = "NoFill"
        except:
            pass
        try:
            if s.line.color.type == 1:
                line = f"RGB({s.line.color.rgb})"
        except:
            pass
        print(f"[{s_idx}] {s.name} ({s.shape_type}): left={s.left}, top={s.top}, w={s.width}, h={s.height}, fill={fill}, line={line}, text=\"{t}\"")
        if s.shape_type == 6:
            for c_idx, c in enumerate(s.shapes):
                ct = c.text.replace('\n', ' // ') if c.has_text_frame else ''
                c_fill = "None"
                c_line = "None"
                try:
                    if c.fill.type == 1: c_fill = f"RGB({c.fill.fore_color.rgb})"
                except: pass
                try:
                    if c.line.color.type == 1: c_line = f"RGB({c.line.color.rgb})"
                except: pass
                print(f"    Child [{c_idx}] {c.name}: left={c.left}, top={c.top}, w={c.width}, h={c.height}, fill={c_fill}, line={c_line}, text=\"{ct}\"")
