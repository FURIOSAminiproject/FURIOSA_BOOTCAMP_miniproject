# -*- coding: utf-8 -*-
"""
apply_ppt_updates.py
퓨리오사AIagent스쿨_미니프로젝트_IamyourJ_final_이지석_장동율.pptx의
Slide 0(표지 이름 볼드), Slide 7(여행 경로 생성), Slide 15(CI/CD) 전문 엔지니어링 콘텐츠 반영
"""

import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import os
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor

PPT_PATH = r"c:\furiosa-miniproject\FURIOSA_BOOTCAMP_miniproject\퓨리오사AIagent스쿨_미니프로젝트_IamyourJ_final_이지석_장동율.pptx"
IMG_SLIDE_7 = r"c:\furiosa-miniproject\FURIOSA_BOOTCAMP_miniproject\slide_07_body.png"
IMG_SLIDE_15 = r"c:\furiosa-miniproject\FURIOSA_BOOTCAMP_miniproject\slide_15_body.png"

prs = Presentation(PPT_PATH)

# =============================================================================
# 1. Slide 0: 표지 저자명 장동율 볼드 처리
# =============================================================================
slide0 = prs.slides[0]
for shape in slide0.shapes:
    if shape.has_text_frame and '장동율' in shape.text:
        print(f"[Slide 0] Found author shape: {shape.name}")
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                if '장동율' in r.text:
                    r.font.bold = True
                    print(f"  -> Updated '장동율' font.bold = True")

# =============================================================================
# 2. Slide 7: BackEnd – 여행 이동 경로 최적화 & 길목 명소 큐레이션
# =============================================================================
slide7 = prs.slides[7]
print("\n[Slide 7] Processing shapes...")
shapes_to_remove = []

for shape in slide7.shapes:
    if shape.has_text_frame:
        text = shape.text.strip()
        if "여행 경로" in text:
            print(f"  Updating title shape: {shape.name}")
            shape.width = 9200000
            shape.text_frame.text = ""
            p = shape.text_frame.paragraphs[0]
            p.text = "BackEnd – 여행 이동 경로 최적화 & 길목 명소 큐레이션"
            p.font.name = "Malgun Gothic"
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = RGBColor(10, 57, 129) # #0A3981
        elif "사용한 API" in text or "여행 경로 생성할때" in text:
            print(f"  Marking placeholder shape for removal: {shape.name}")
            shapes_to_remove.append(shape)

for shape in shapes_to_remove:
    sp = shape._element
    sp.getparent().remove(sp)
    print(f"  -> Removed shape {shape.name}")

# 본문 그래픽 삽입
print(f"  Adding picture {IMG_SLIDE_7} to Slide 7...")
slide7.shapes.add_picture(
    IMG_SLIDE_7,
    left=320000,
    top=1220000,
    width=11552000,
    height=5408000
)
print("  -> Picture added successfully to Slide 7")

# =============================================================================
# 3. Slide 15: CI/CD – GitHub Actions & Docker Compose 무중단 자동 배포 파이프라인
# =============================================================================
slide15 = prs.slides[15]
print("\n[Slide 15] Processing shapes...")
shapes_to_remove = []

for shape in slide15.shapes:
    if shape.has_text_frame:
        text = shape.text.strip()
        if "CI/CD" in text and "어떤거" not in text:
            print(f"  Updating title shape: {shape.name}")
            shape.width = 9800000
            shape.text_frame.text = ""
            p = shape.text_frame.paragraphs[0]
            p.text = "CI/CD – GitHub Actions & Docker 무중단 자동 배포 파이프라인"
            p.font.name = "Malgun Gothic"
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = RGBColor(10, 57, 129) # #0A3981
        elif "어떤거" in text or "자동 배포" in text or "피드백" in text:
            print(f"  Marking placeholder shape for removal: {shape.name}")
            shapes_to_remove.append(shape)

for shape in shapes_to_remove:
    sp = shape._element
    sp.getparent().remove(sp)
    print(f"  -> Removed shape {shape.name}")

# 본문 그래픽 삽입
print(f"  Adding picture {IMG_SLIDE_15} to Slide 15...")
slide15.shapes.add_picture(
    IMG_SLIDE_15,
    left=320000,
    top=1220000,
    width=11552000,
    height=5408000
)
print("  -> Picture added successfully to Slide 15")

# 저장
prs.save(PPT_PATH)
print(f"\n[성공] 프레젠테이션 업데이트 및 저장 완료: {PPT_PATH}")
