from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.oxml import parse_xml
from pptx.util import Pt
from pptx.dml.color import RGBColor

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[6])
conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 1000000, 2000000, 2000000, 2000000)
conn.line.color.rgb = RGBColor(13, 148, 136)
conn.line.width = Pt(2.5)
tail = parse_xml('<a:tailEnd xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" type="triangle" w="med" len="med"/>')
conn.line._ln.append(tail)
print("Success! XML:", conn._element.xml[:200])
