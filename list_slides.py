import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pptx import Presentation

ppt_path = r"c:\furiosa-miniproject\FURIOSA_BOOTCAMP_miniproject\퓨리오사AIagent스쿨_미니프로젝트_IamyourJ_final_이지석_장동율.pptx"
prs = Presentation(ppt_path)

for idx, slide in enumerate(prs.slides):
    texts = []
    for s in slide.shapes:
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                txt = p.text.strip()
                if txt:
                    texts.append(txt)
    title = texts[0] if texts else "No text"
    sub = texts[1] if len(texts) > 1 else ""
    print(f"Slide {idx} (page {idx}): {title} | {sub[:40]}")
