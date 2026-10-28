# -*- coding: utf-8 -*-
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
import os, shutil

prs = Presentation(r"c:\furiosa-miniproject\FURIOSA_BOOTCAMP_miniproject\퓨리오사AIagent스쿨_미니프로젝트_IamyourJ_final_이지석_장동율.pptx")

SCALE_X = 1920.0 / 12192000.0
SCALE_Y = 1080.0 / 6858000.0

try:
    font_title = ImageFont.truetype("malgunbd.ttf", 36)
    font_sub = ImageFont.truetype("malgunbd.ttf", 22)
    font_num = ImageFont.truetype("malgun.ttf", 24)
    font_box_lg = ImageFont.truetype("malgunbd.ttf", 20)
    font_box_md = ImageFont.truetype("malgunbd.ttf", 16)
    font_box_sm = ImageFont.truetype("malgun.ttf", 14)
    font_code = ImageFont.truetype("consola.ttf", 14)
except:
    font_title = font_sub = font_num = font_box_lg = font_box_md = font_box_sm = font_code = ImageFont.load_default()

def sanitize(txt):
    return str(txt).replace('\r', '').replace('\n', '').replace('\x0b', '').strip()

def render_slide(slide, out_path):
    im = Image.new("RGB", (1920, 1080), "#FFFFFF")
    draw = ImageDraw.Draw(im)
    draw.line([(0, int(1136073 * SCALE_Y)), (1920, int(1136073 * SCALE_Y))], fill="#223B6B", width=2)
    
    for s in slide.shapes:
        sx = int(s.left * SCALE_X)
        sy = int(s.top * SCALE_Y)
        sw = int(s.width * SCALE_X)
        sh = int(s.height * SCALE_Y)
        
        if s.name == "TextBox 2" or s.name == "TextBox 3" or s.name == "TextBox 48" or (s.shape_type == 17 and sy < 100):
            t = sanitize(s.text_frame.paragraphs[0].text) if s.has_text_frame and s.text_frame.paragraphs else sanitize(s.text)
            draw.text((sx, sy), t, fill="#223B6B", font=font_title)
        elif s.name == "슬라이드 번호 개체 틀 1" or (s.shape_type == 17 and sx > 1400 and sy < 100):
            t = sanitize(s.text)
            draw.text((sx + sw - 60, sy), t, fill="#475569", font=font_num)
        elif "TextBox" in s.name or s.shape_type == 17:
            if "Multi-stage" in s.text or "Two-stage" in s.text or "OpenRouteService" in s.text or "공간 점-선분" in s.text:
                t = sanitize(s.text_frame.paragraphs[0].text) if s.has_text_frame and s.text_frame.paragraphs else sanitize(s.text)
                draw.text((sx, sy), t, fill="#475569", font=font_sub)
            else:
                if s.has_text_frame:
                    paras = [sanitize(p.text) for p in s.text_frame.paragraphs if sanitize(p.text)]
                    ly = sy + 6
                    for idx, line in enumerate(paras):
                        f = font_box_lg if idx == 0 else font_box_md
                        c = "#0F172A" if idx == 0 else "#475569"
                        draw.text((sx + sw // 2, ly), line, fill=c, font=f, anchor="mm")
                        ly += 26
        elif s.shape_type == 1: # Rounded rectangle
            fill_c = "#FFFFFF"
            border_c = "#94A3B8"
            try:
                if s.fill.type == 1:
                    rgb = s.fill.fore_color.rgb
                    fill_c = f"#{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"
            except: pass
            try:
                if s.line.color.type == 1:
                    rgb = s.line.color.rgb
                    border_c = f"#{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"
            except: pass
            
            draw.rounded_rectangle([sx, sy, sx + sw, sy + sh], radius=10, fill=fill_c, outline=border_c, width=2)
            
            if s.has_text_frame:
                paras = [sanitize(p.text) for p in s.text_frame.paragraphs if sanitize(p.text)]
                n_lines = len(paras)
                align_left = False
                try:
                    if s.text_frame.paragraphs[0].alignment == 1: # LEFT
                        align_left = True
                except: pass
                
                if n_lines > 0:
                    total_h = n_lines * 22
                    cur_y = sy + max(14, (sh - total_h) // 2)
                    for p in s.text_frame.paragraphs:
                        txt = sanitize(p.text)
                        if not txt:
                            cur_y += 8
                            continue
                        f_to_use = font_box_lg if p.font.size and p.font.size.pt >= 15 else (font_box_md if p.font.size and p.font.size.pt >= 13 else font_box_sm)
                        col = "#0F172A"
                        try:
                            if p.font.color and p.font.color.rgb:
                                rgb = p.font.color.rgb
                                col = f"#{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"
                        except: pass
                        if align_left:
                            draw.text((sx + 16, cur_y), txt, fill=col, font=f_to_use)
                        else:
                            draw.text((sx + sw // 2, cur_y + 10), txt, fill=col, font=f_to_use, anchor="mm")
                        cur_y += 24
                        
        elif s.shape_type == 9: # Connector / Line
            if s.width > 0 and s.height == 0 and s.top > 1500000:
                line_col = "#0D9488"
                draw.line([(sx, sy), (sx + sw, sy)], fill=line_col, width=4)
                ah_w = 14
                ah_h = 9
                draw.polygon([
                    (sx + sw, sy),
                    (sx + sw - ah_w, sy - ah_h),
                    (sx + sw - ah_w, sy + ah_h)
                ], fill=line_col)
            elif (s.width == 0 or s.width < 100000) and s.height > 1000000:
                draw.line([(sx, sy), (sx, sy + sh)], fill="#CBD5E1", width=2)
            
        elif s.name == "Rectangle 1" or (s.width < 100000 and s.height > 1000000):
            draw.line([(sx, sy), (sx, sy + sh)], fill="#CBD5E1", width=2)

    im.save(out_path, quality=95)
    print("Rendered:", out_path)

art_dir = r"C:\Users\admin\.gemini\antigravity-ide\brain\bd09dd9a-d208-4cbc-a866-27dcdcfc298c"

for s_idx in [7, 8, 9, 10, 18]:
    out_name = f"preview_final_slide_{s_idx:02d}.png"
    out_p = os.path.join(art_dir, out_name)
    render_slide(prs.slides[s_idx], out_p)

print("All slide previews generated cleanly!")
