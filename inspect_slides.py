import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pptx import Presentation

ppt_path = r"c:\furiosa-miniproject\FURIOSA_BOOTCAMP_miniproject\퓨리오사AIagent스쿨_미니프로젝트_IamyourJ_final_이지석_장동율.pptx"
prs = Presentation(ppt_path)

for slide_idx in [7, 12, 15]:
    slide = prs.slides[slide_idx]
    print(f"=== Slide {slide_idx} ===")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            tf = shape.text_frame
            print(f"Shape {s_idx} ({shape.name}):")
            for p_idx, p in enumerate(tf.paragraphs):
                runs = [r.text for r in p.runs]
                print(f"  P{p_idx}: text={repr(p.text)} (chars={[hex(ord(c)) for c in p.text]}) runs={runs}")
